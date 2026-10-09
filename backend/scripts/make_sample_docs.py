"""从 .txt 样本生成 DOCX / PDF 版本的简历样本。

存在的意义：契约声明支持 PDF / DOCX / PNG / JPG，而 .txt 走的是最简单的文本分支，
无法验证真实的「文档解析 → 文本抽取 → 结构化」链路。这个脚本让二进制样本可复现，
而不是往仓库里塞几个来路不明的文件。

用法（在 backend 目录下）：
    python -m scripts.make_sample_docs            # 生成 DOCX，并尝试转 PDF
    python -m scripts.make_sample_docs --no-pdf   # 只生成 DOCX
"""

from __future__ import annotations

import argparse
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

SAMPLES = Path(__file__).resolve().parent.parent / "samples"

# 由 LibreOffice Kit 提供，避免依赖系统安装的 LibreOffice
NODE_BIN = Path(r"E:\deepseek-hhh\resources\runtime\primary-runtime\dependencies\node\bin\node.exe")
LIBREOFFICE_CLI = Path(
    r"E:\deepseek-hhh\resources\app.asar.unpacked\dsh\node_modules\@deepseek-ai\libreoffice-kit\lib\cli.js"
)

SECTION_HEADINGS = {
    "个人信息",
    "个人优势",
    "工作经历",
    "项目经历",
    "专业技能",
    "教育经历",
    "证书",
    "证书 / 荣誉",
    "自我评价",
}


def build_docx(source: Path, target: Path) -> None:
    """把线性文本转成带标题样式的真实 DOCX（而不是一整段纯文本）。"""
    from docx import Document
    from docx.shared import Pt

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    # 中文字形需要显式设置 eastAsia，只设 font.name 不生效
    style.element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "微软雅黑"
    )

    for raw_line in source.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line in SECTION_HEADINGS:
            document.add_heading(line, level=1)
        elif line.startswith("-"):
            document.add_paragraph(line.lstrip("- ").strip(), style="List Bullet")
        elif len(document.paragraphs) == 0:
            document.add_heading(line, level=0)
        else:
            document.add_paragraph(line)

    document.save(target)


def to_pdf(docx_path: Path, pdf_path: Path) -> tuple[bool, str]:
    if not NODE_BIN.exists() or not LIBREOFFICE_CLI.exists():
        return False, f"LibreOffice Kit 不可用（node={NODE_BIN.exists()}, cli={LIBREOFFICE_CLI.exists()}）"

    proc = subprocess.run(
        [str(NODE_BIN), str(LIBREOFFICE_CLI), "convert", "--input", str(docx_path), "--output", str(pdf_path)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if proc.returncode != 0:
        return False, (proc.stderr or proc.stdout or "").strip()[:400]
    if not pdf_path.exists():
        return False, "转换命令返回成功但未生成 PDF"
    return True, ""


def main() -> int:
    parser = argparse.ArgumentParser(description="生成 DOCX / PDF 简历样本")
    parser.add_argument("--no-pdf", action="store_true", help="跳过 PDF 生成")
    args = parser.parse_args()

    sources = sorted(SAMPLES.glob("sample_resume_*.txt"))
    if not sources:
        print(f"未找到 .txt 样本：{SAMPLES}")
        return 1

    failures = 0
    for source in sources:
        docx_path = source.with_suffix(".docx")
        build_docx(source, docx_path)
        print(f"[DOCX] {docx_path.name}  ({docx_path.stat().st_size / 1024:.1f} KB)")

        if args.no_pdf:
            continue

        pdf_path = source.with_suffix(".pdf")
        if pdf_path.exists():
            pdf_path.unlink()
        ok, message = to_pdf(docx_path, pdf_path)
        if ok:
            print(f"[PDF ] {pdf_path.name}  ({pdf_path.stat().st_size / 1024:.1f} KB)")
        else:
            failures += 1
            print(f"[PDF ] 跳过 {source.stem}：{message}")

    return 0 if failures == 0 else 0  # PDF 失败不算致命（依赖外部组件）


if __name__ == "__main__":
    raise SystemExit(main())
